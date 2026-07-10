"""Generate a mock Amazon India jewellery flat-file template.

The real template must be downloaded from Seller Central; this script
produces a structurally faithful stand-in (banner row, display-name row,
field-name header row, drop-down data validation) for demos and tests.

Usage:
    python sample_data/make_sample_template.py [output.xlsx]
"""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

FIELDS: list[tuple[str, str]] = [
    ("TemplateType=Jewelry", "feed_product_type"),
    ("Seller SKU", "item_sku"),
    ("Brand Name", "brand_name"),
    ("Product Name", "item_name"),
    ("Manufacturer", "manufacturer"),
    ("Part Number", "part_number"),
    ("Product ID", "external_product_id"),
    ("Product ID Type", "external_product_id_type"),
    ("Recommended Browse Nodes", "recommended_browse_nodes"),
    ("Standard Price", "standard_price"),
    ("MRP", "maximum_retail_price"),
    ("Quantity", "quantity"),
    ("Add/Delete", "update_delete"),
    ("Product Description", "product_description"),
    ("Bullet Point 1", "bullet_point1"),
    ("Bullet Point 2", "bullet_point2"),
    ("Bullet Point 3", "bullet_point3"),
    ("Bullet Point 4", "bullet_point4"),
    ("Bullet Point 5", "bullet_point5"),
    ("Search Terms", "generic_keywords"),
    ("Main Image URL", "main_image_url"),
    ("Other Image URL 1", "other_image_url1"),
    ("Other Image URL 2", "other_image_url2"),
    ("Other Image URL 3", "other_image_url3"),
    ("Other Image URL 4", "other_image_url4"),
    ("Other Image URL 5", "other_image_url5"),
    ("Other Image URL 6", "other_image_url6"),
    ("Other Image URL 7", "other_image_url7"),
    ("Other Image URL 8", "other_image_url8"),
    ("Parentage", "parent_child"),
    ("Parent SKU", "parent_sku"),
    ("Relationship Type", "relationship_type"),
    ("Variation Theme", "variation_theme"),
    ("Country of Origin", "country_of_origin"),
    ("Metal Type", "metal_type"),
    ("Metal Stamp", "metal_stamp"),
    ("Stone Creation Method", "stone_creation_method"),
    ("Stone Type", "stone_type"),
    ("Stone Shape", "stone_shape"),
    ("Stone Weight", "stone_weight"),
    ("Stone Clarity", "stone_clarity"),
    ("Ring Size", "ring_size"),
    ("Item Weight", "item_weight"),
    ("Item Weight Unit", "item_weight_unit_of_measure"),
    ("Department", "department_name"),
    ("Target Gender", "target_gender"),
    ("Handling Time", "fulfillment_latency"),
    ("Colour", "color_name"),
    ("Size", "size_name"),
    ("Material Type", "material_type"),
]


def build_template(path: Path) -> Path:
    """Create the mock template workbook at ``path``."""
    workbook = Workbook()

    instructions = workbook.active
    instructions.title = "Instructions"
    instructions["A1"] = "Mock Amazon India Jewellery template for testing."

    valid_values = workbook.create_sheet("Valid Values")
    valid_values["A1"] = "parent_child"
    valid_values["A2"] = "Parent"
    valid_values["A3"] = "Child"

    sheet = workbook.create_sheet("Template")
    sheet["A1"] = "TemplateType=Jewelry"
    sheet["B1"] = "Version=2026.0710"
    banner_font = Font(bold=True, color="FFFFFF")
    banner_fill = PatternFill("solid", fgColor="4472C4")
    for col, (display, field) in enumerate(FIELDS, start=1):
        display_cell = sheet.cell(row=2, column=col, value=display)
        display_cell.font = banner_font
        display_cell.fill = banner_fill
        sheet.cell(row=3, column=col, value=field)

    parentage_col = next(
        i for i, (_, f) in enumerate(FIELDS, start=1) if f == "parent_child"
    )
    dv = DataValidation(type="list", formula1='"parent,child"', allow_blank=True)
    sheet.add_data_validation(dv)
    from openpyxl.utils import get_column_letter

    letter = get_column_letter(parentage_col)
    dv.add(f"{letter}4:{letter}20000")

    sheet.freeze_panes = "A4"
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


if __name__ == "__main__":
    default = Path(__file__).resolve().parent.parent / "templates" / "Amazon_Template_Sample.xlsx"
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else default
    print(f"Wrote {build_template(target)}")
