"""Tests for template header detection and the Excel writer."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from core.amazon_template import AmazonTemplate
from core.excel_writer import ExcelWriter
from core.models import MappingRule, SummaryStats
from core.utils import TemplateError


def test_header_row_detected_dynamically(sample_template: Path) -> None:
    layout = AmazonTemplate(sample_template).layout()
    # The mock template has a banner row and display-name row above the
    # real field-name row.
    assert layout.header_row == 3
    assert layout.data_start_row == 4
    assert layout.sheet_name == "Template"


def test_columns_found_by_name_not_position(sample_template: Path) -> None:
    layout = AmazonTemplate(sample_template).layout()
    assert layout.column_for("item_sku") is not None
    assert layout.column_for("ITEM-SKU") == layout.column_for("item_sku")
    assert layout.column_for("does_not_exist") is None


def test_missing_template_raises(tmp_path: Path) -> None:
    with pytest.raises(TemplateError, match="not found"):
        AmazonTemplate(tmp_path / "missing.xlsm").load()


def test_writer_populates_below_header(sample_template: Path, tmp_path: Path) -> None:
    template = AmazonTemplate(sample_template)
    amazon = pd.DataFrame(
        {
            "item_sku": ["SKU-1", "SKU-2"],
            "item_name": ["Ring A", "Ring B"],
            "not_in_template_field": ["x", "y"],
        }
    )
    rules = [
        MappingRule(
            amazon_field="item_sku",
            aliases=["seller-sku"],
            source={"type": "internal", "column": "_sku"},
        ),
        MappingRule(
            amazon_field="item_name",
            source={"type": "internal", "column": "_seo_title"},
        ),
        MappingRule(
            amazon_field="not_in_template_field",
            source={"type": "constant", "value": "x"},
        ),
    ]
    out_path, skipped = ExcelWriter().write_upload_file(
        template, amazon, rules, tmp_path / "upload.xlsx"
    )
    assert skipped == ["not_in_template_field"]

    workbook = load_workbook(out_path)
    sheet = workbook["Template"]
    layout = AmazonTemplate(out_path).layout()
    sku_col = layout.column_for("item_sku")
    assert sheet.cell(row=4, column=sku_col).value == "SKU-1"
    assert sheet.cell(row=5, column=sku_col).value == "SKU-2"
    # Header rows untouched.
    assert sheet.cell(row=3, column=sku_col).value == "item_sku"
    # Data validation (drop-downs) preserved.
    assert len(sheet.data_validations.dataValidation) == 1


def test_issue_report_written(tmp_path: Path) -> None:
    issues = pd.DataFrame(
        {
            "Row": [1, 2],
            "SKU": ["A", "B"],
            "Field": ["item_sku", "standard_price"],
            "Severity": ["Error", "Warning"],
            "Message": ["Missing SKU", "Missing product weight"],
        }
    )
    path = ExcelWriter().write_issue_report(issues, tmp_path / "v.xlsx", "Validation")
    sheet = load_workbook(path).active
    assert sheet.max_row == 3  # header + 2 findings
    error_path = ExcelWriter().write_issue_report(
        issues, tmp_path / "e.xlsx", "Errors", errors_only=True
    )
    assert load_workbook(error_path).active.max_row == 2  # header + 1 error


def test_summary_report_written(tmp_path: Path) -> None:
    stats = SummaryStats(total_products=4, parents=2, children=6, errors=1)
    path = ExcelWriter().write_summary_report(stats, tmp_path / "s.xlsx")
    sheet = load_workbook(path).active
    metrics = {row[0].value: row[1].value for row in sheet.iter_rows(min_row=2)}
    assert metrics["Total Products"] == 4
    assert metrics["Parent Listings"] == 2
    assert metrics["Validation Errors"] == 1
